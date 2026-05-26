"""Corpora_Meta_Table.xlsx (RAW_통합 시트) → data/GOLD/corpora.jsonl 변환.

사용:
    python scripts/build_corpus_meta.py \\
        --in data/Corpora_Meta_Table.xlsx \\
        --out data/GOLD/corpora.jsonl \\
        --only-marked         # 'o' 표시된 행만 (선택)

원칙 (docs/data/schema.md §6 자동화 정책):
- 수동 변환 금지 — 이 스크립트가 단일 진실
- 컬럼 매핑은 본 파일 상단 MAPPING dict 에서 한 곳 관리
- 미지의 컬럼은 extra 로 보존
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl 가 필요합니다. `pip install openpyxl`", file=sys.stderr)
    sys.exit(1)


# RAW_통합 시트 컬럼 → CorpusMeta 필드 매핑
# (None 은 드롭, 그 외는 키로 변환)
MAPPING: dict[str, str | None] = {
    "D": None,                       # 'o' 표시 컬럼 (자체 메타 아님)
    "Title (U)":     "corpus_id",
    "Desc":          "title",
    "Domain":        "domain_hint",
    "Style":         "style",
    "Language":      "language",
    "File Unit":     None,
    "# Utts":        None,           # 코퍼스 단위 통계는 시트 그대로
    "Hours":         None,
    "# Speakers":    None,
    "File Size":     None,
    "Sample\nRate":  None,
    "Bit\nDepth":    None,
    "Channels":      None,
    "Has\nTranscript": None,
    "Speaker\nInfo": None,
    "Speaker\nNote": "speaker_note",
    "Recording\nEnv":"recording_env",
    "Emotion/\nStyle Label": None,
    "Source":        "source",
    "Link":          "url",
    "License":       "license",
    "Version":       "version",
    "Tag":           None,
    "수집일\n(yyyy-mm)": "collected_at",
    "비고":          None,
}

MARK_COL = "D"        # 'o' 표시 컬럼 (또는 col1)
SHEET_NAME = "RAW_통합"


def load_rows(xlsx_path: Path, *, only_marked: bool = False) -> list[dict]:
    """xlsx → list of dict. 헤더는 1행."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(
            f"시트 '{SHEET_NAME}' 을 찾을 수 없음. 가능한 시트: {wb.sheetnames}"
        )
    ws = wb[SHEET_NAME]
    headers: list[str] = []
    for c in ws[1]:
        v = c.value
        if v is None:
            headers.append("D")     # 첫 컬럼 라벨이 비어있을 수 있음
        else:
            headers.append(str(v))

    rows: list[dict] = []
    for r in range(2, ws.max_row + 1):
        row = {h: ws.cell(r, i + 1).value for i, h in enumerate(headers)}
        # 빈 행 스킵
        if not row.get("Title (U)"):
            continue
        if only_marked and row.get(MARK_COL) != "o":
            continue
        rows.append(row)
    return rows


def convert_row(row: dict) -> dict:
    """xlsx 한 행 → JSONL 한 줄 (dict)."""
    out: dict = {}
    extra: dict = {}
    for col, val in row.items():
        if val is None:
            continue
        if col in MAPPING:
            tgt = MAPPING[col]
            if tgt is None:
                continue
            out[tgt] = val
        else:
            # 매핑에 없는 컬럼 → 사용자가 새로 추가했을 수 있음. extra 로 보존.
            extra[col] = val

    # 기본값 보강
    out.setdefault("title", out.get("corpus_id", ""))
    out.setdefault("language", "ko")

    # 타입 정리 (datetime 등)
    for k, v in list(out.items()):
        if hasattr(v, "isoformat"):
            out[k] = v.isoformat()

    if extra:
        out["extra"] = {k: (v.isoformat() if hasattr(v, "isoformat") else v)
                        for k, v in extra.items()}
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--in", dest="in_xlsx",
                    default="data/Corpora_Meta_Table.xlsx")
    ap.add_argument("--out", dest="out_jsonl",
                    default="data/GOLD/corpora.jsonl")
    ap.add_argument("--only-marked", action="store_true",
                    help="첫 컬럼에 'o' 표시된 행만 출력")
    args = ap.parse_args()

    in_path = Path(args.in_xlsx)
    out_path = Path(args.out_jsonl)
    if not in_path.exists():
        raise SystemExit(f"입력 xlsx 없음: {in_path}")

    rows = load_rows(in_path, only_marked=args.only_marked)
    if not rows:
        raise SystemExit("변환할 행이 없습니다.")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    written = 0
    seen_ids: set[str] = set()
    with out_path.open("w", encoding="utf-8") as f:
        for row in rows:
            d = convert_row(row)
            cid = d.get("corpus_id")
            if not cid:
                continue
            if cid in seen_ids:
                print(f"WARN: 중복 corpus_id 스킵: {cid!r}", file=sys.stderr)
                continue
            seen_ids.add(cid)
            f.write(json.dumps(d, ensure_ascii=False) + "\n")
            written += 1

    print(f"wrote {written} rows → {out_path}")
    if args.only_marked:
        print(f"  (only-marked: {SHEET_NAME}.{MARK_COL} == 'o')")


if __name__ == "__main__":
    main()
