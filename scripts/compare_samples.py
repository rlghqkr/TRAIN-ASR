"""발화(샘플) 단위 결과 비교 → 엑셀 (CLI).

같은 벤치마크를 여러 모델로 평가한 `predictions.jsonl` 들을 **발화 키(key)로 join**
해서, "한 발화를 각 모델이 어떻게 인식했나"를 나란히 보는 엑셀을 만든다.
(모델 집계 비교는 compare_results.py, 이건 그보다 아래 층위인 발화별 비교.)

⚠ 실행 환경: jiwer 가 필요하므로 학습/평가와 같은 conda env 에서 돌린다.
    conda run -n train-asr python scripts/compare_samples.py ...
    (또는  ~/anaconda3/envs/train-asr/bin/python scripts/compare_samples.py ...)

무엇을 만드나 (엑셀 시트):
    - 요약        : 모델별 CER/sCER/WER + 완전정답 수 + 공백만틀린 수. 2모델이면 해설.
    - 문장별비교  : key·duration·정답 + [모델별 예측·CER] + Δ(2모델) + winner + raw예측.
    - <모델>_치환/삽입/삭제 : 단어단위 오류 집계(정답↔예측). 어디서 자주 틀리는지.

입력(모델 1개당): 결과 폴더 또는 모델 이름 (BENCHMARK/results/<name>/ 로 해석).
    그 폴더 안 <benchmark>/predictions.jsonl 을 읽는다.

사용법:
    # kiosk 벤치를 세 모델로 발화 비교
    ENV=~/anaconda3/envs/train-asr/bin/python
    $ENV scripts/compare_samples.py -b kiosk_cafe_order \
        whisper_medium_baseline_0701 \
        whisper_medium_baseline_0701_adapt_cafe__260702_033807 \
        whisper_large_v3_baseline \
        --labels medium,adapt_cafe,large_v3

    # 틀린 발화만(하나라도 CER>0), 오답 큰 순으로 정렬
    $ENV scripts/compare_samples.py -b kiosk_cafe_order A B --only-diff --sort worst
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

import pandas as pd
import structlog
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# 프로젝트 루트를 import 경로에 추가 (project.* 재사용).
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from project.evaluation.cer import compute_cer  # noqa: E402

try:
    import jiwer
except ImportError as e:                                    # 친절한 안내(조용한 실패 금지)
    raise ImportError(
        "jiwer 가 필요합니다. 학습/평가 env 에서 실행하세요: "
        "`~/anaconda3/envs/train-asr/bin/python scripts/compare_samples.py ...`"
    ) from e


logger = structlog.get_logger()

RESULTS_DIR = ROOT / "BENCHMARK" / "results"
PRED_NAME = "predictions.jsonl"
REF_FIELD = "text_norm"                 # 정답(이미 정규화됨 — 리포트와 동일 기준)
HYP_FIELD = "prediction_normalized"     # 예측(정규화). CER 은 이 둘로 계산.
RAW_FIELD = "prediction_raw"            # 예측 원문(정규화 전) — 참고 열.

# 폴더명 끝 실행 타임스탬프( __YYMMDD_HHMMSS ) — 짧은 열이름 만들 때 제거.
_TIMESTAMP_RE = re.compile(r"__\d{6}_\d{6}$")

# 엑셀 스타일.
_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_WIN_FILL = PatternFill("solid", fgColor="C6EFCE")         # 승자(최저 CER) 초록
_WIN_FONT = Font(color="006100")
_ERR_FILL = PatternFill("solid", fgColor="FFF2CC")         # 완전오답(CER=100) 연노랑
_FMT2 = "0.00"


# ───────────────────────────────────────────────────────────────────────
# 1. 입력 해석 & 로딩
# ───────────────────────────────────────────────────────────────────────
def resolve_model_dir(arg: str) -> Path:
    """모델 인자 → 결과 폴더 경로 (폴더 직접 지정 또는 RESULTS_DIR/<이름>)."""
    p = Path(arg)
    if p.is_dir():
        return p
    cand = RESULTS_DIR / arg
    if cand.is_dir():
        return cand
    raise FileNotFoundError(
        f"'{arg}' 결과 폴더를 못 찾음. 시도: {p} , {cand}"
    )


def load_predictions(model_dir: Path, benchmark: str) -> dict[str, dict]:
    """<model_dir>/<benchmark>/predictions.jsonl → {key: sample}.

    Raises:
        FileNotFoundError: 해당 벤치의 predictions.jsonl 이 없을 때(그 모델은 이
            벤치를 안 돌린 것 — 조용히 넘기지 않고 명시적 에러).
    """
    path = model_dir / benchmark / PRED_NAME
    if not path.is_file():
        raise FileNotFoundError(
            f"{path} 없음 — '{model_dir.name}' 는 '{benchmark}' 를 평가하지 않았음"
        )
    out: dict[str, dict] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            d = json.loads(line)
            out[d["key"]] = d
    return out


def make_labels(model_dirs: list[Path], override: list[str] | None) -> list[str]:
    """열(모델) 이름 — 기본은 폴더명에서 타임스탬프 제거, 겹치면 전체명 유지."""
    if override is not None:
        if len(override) != len(model_dirs):
            raise ValueError(
                f"--labels 개수({len(override)}) != 모델 개수({len(model_dirs)})"
            )
        return override
    full = [d.name for d in model_dirs]
    short = [_TIMESTAMP_RE.sub("", n) for n in full]
    dupes = {s for s in short if short.count(s) > 1}
    return [full[i] if short[i] in dupes else short[i] for i in range(len(full))]


# ───────────────────────────────────────────────────────────────────────
# 2. 발화별 CER / sCER
# ───────────────────────────────────────────────────────────────────────
# 정답·예측 모두 normalize_korean_asr 로 구두점이 이미 제거된 상태(text_norm/
# prediction_normalized). 여기에:
#   CER  = 공백 포함 (구두점 제외) — 띄어쓰기 오류까지 반영
#   sCER = 공백 무시 (구두점+띄어쓰기 제외) — 순수 글자 정확도. Whisper 스페이싱이
#          불안정하므로 "내용을 맞췄나"는 sCER 이 공정. (cer.py 의 scer 와 동일 정의)
def sample_cer(ref: str, hyp: str) -> float | None:
    """발화 1개 CER(%, 공백 포함). 빈 정답은 None."""
    if not ref.strip():
        return None
    return round(jiwer.cer(ref, hyp) * 100, 2)


def sample_scer(ref: str, hyp: str) -> float | None:
    """발화 1개 sCER(%, 구두점+띄어쓰기 제외). 빈 정답은 None."""
    if not ref.strip():
        return None
    return round(jiwer.cer(ref.replace(" ", ""), hyp.replace(" ", "")) * 100, 2)


# 표시 이름 → 계산 함수. --metric 이 이 중 무엇을 낼지 고른다.
METRIC_FNS: dict[str, "callable"] = {"CER": sample_cer, "sCER": sample_scer}


# ───────────────────────────────────────────────────────────────────────
# 3. 요약 시트
# ───────────────────────────────────────────────────────────────────────
def build_summary(
    preds: list[dict[str, dict]], labels: list[str],
) -> pd.DataFrame:
    """모델별 코퍼스 CER/sCER/WER + 완전정답 수 + 공백만틀린 수. 2모델이면 해설."""
    stats: list[dict] = []
    for data in preds:
        refs = [s[REF_FIELD] for s in data.values()]
        hyps = [s[HYP_FIELD] for s in data.values()]
        r = compute_cer(refs, hyps, normalize=False)        # 리포트와 동일 기준
        perfect = sum(1 for a, b in zip(refs, hyps)
                      if a.strip() and a == b)
        space_only = sum(                                    # 공백만 틀림
            1 for a, b in zip(refs, hyps)
            if a != b and a.replace(" ", "") == b.replace(" ", "")
        )
        stats.append({
            "utterances": r.samples,
            "CER (공백포함) %": round(r.cer, 3),
            "sCER (공백무시) %": round(r.scer, 3),
            "WER %": round(r.wer, 3),
            "완전정답(CER=0) 수": perfect,
            "공백만틀린 수": space_only,
        })

    rows: list[dict] = []
    for metric in stats[0]:
        row = {"지표": metric}
        for lab, st in zip(labels, stats):
            row[lab] = st[metric]
        if len(labels) == 2:
            row["해설"] = _explain(metric, stats[0], stats[1], labels)
        rows.append(row)
    return pd.DataFrame(rows)


def _explain(metric: str, a: dict, b: dict, labels: list[str]) -> str:
    """2모델 한정 자동 해설 — 지표별 우세 모델/격차."""
    va, vb = a[metric], b[metric]
    if metric == "utterances":
        return "동일" if va == vb else f"서로 다름 ({va} vs {vb})"
    diff = round(abs(va - vb), 3)
    if va == vb:
        return "동일"
    lower_better = metric != "완전정답(CER=0) 수"             # 완전정답만 높을수록 좋음
    if lower_better:
        win = labels[0] if va < vb else labels[1]
    else:
        win = labels[0] if va > vb else labels[1]
    unit = "개" if metric.endswith("수") else "p"
    return f"{win} 우세 ({diff}{unit})"


# ───────────────────────────────────────────────────────────────────────
# 4. 문장별비교 시트
# ───────────────────────────────────────────────────────────────────────
def build_sentence(
    preds: list[dict[str, dict]], labels: list[str],
    *, metrics: list[str], only_diff: bool, sort: str,
) -> pd.DataFrame:
    """발화별 정답·모델별 예측/CER 나란히. winner·Δ 는 decide 지표(sCER 우선)로.

    Args:
        metrics: 낼 지표 이름 목록(['sCER'] 또는 ['CER','sCER'] 등).
        only_diff: True 면 (decide 지표 기준) 모든 모델이 완전정답인 발화는 제외.
        sort: 'key'(키순) | 'worst'(최대 오류 높은 순) | 'delta'(2모델 차 큰 순).
    """
    all_keys = _union_keys(preds)
    two = len(labels) == 2
    decide = "sCER" if "sCER" in metrics else "CER"          # 승자 판정 지표
    delta_col = f"Δ{decide}({labels[0]}-{labels[1]})" if two else None

    rows: list[dict] = []
    for key in all_keys:
        present = [(lab, data[key]) for lab, data in zip(labels, preds)
                   if key in data]
        ref = present[0][1][REF_FIELD]                       # 정답(모델 공통)
        dur = present[0][1].get("duration")
        # 지표별 {label: 값} 미리 계산.
        scores = {m: {lab: METRIC_FNS[m](ref, s[HYP_FIELD]) for lab, s in present}
                  for m in metrics}
        dvals = {l: c for l, c in scores[decide].items() if c is not None}
        if only_diff and dvals and all(c == 0 for c in dvals.values()):
            continue                                          # 전원 정답 → 스킵

        row: dict = {"key": key, "duration": dur, "ref": ref}
        for lab in labels:
            idx = labels.index(lab)
            row[f"{lab}_pred"] = preds[idx].get(key, {}).get(HYP_FIELD)
            for m in metrics:
                row[f"{lab}_{m}"] = scores[m].get(lab)
        if two:
            ca, cb = scores[decide].get(labels[0]), scores[decide].get(labels[1])
            row[delta_col] = (round(ca - cb, 2)
                              if ca is not None and cb is not None else None)
        row["winner"] = _winner(dvals)
        for lab in labels:                                    # raw 예측(참고, 뒤로)
            row[f"{lab}_raw"] = preds[labels.index(lab)].get(key, {}).get(RAW_FIELD)
        rows.append(row)

    df = pd.DataFrame(rows)
    return _sort_sentence(df, labels, sort, two, decide, delta_col)


def _union_keys(preds: list[dict[str, dict]]) -> list[str]:
    """모든 모델 발화 키 합집합 — 첫 등장 순서 보존(대개 key 정렬과 동일)."""
    seen: list[str] = []
    known: set[str] = set()
    for data in preds:
        for k in data:
            if k not in known:
                known.add(k)
                seen.append(k)
    return seen


def _winner(valid: dict[str, float | None]) -> str:
    """최저 CER 모델 이름. 동률이면 'tie', 없으면 ''."""
    if not valid:
        return ""
    lo = min(valid.values())
    best = [l for l, c in valid.items() if c == lo]
    return best[0] if len(best) == 1 else "tie"


def _sort_sentence(df: pd.DataFrame, labels: list[str], sort: str,
                   two: bool, decide: str, delta_col: str | None) -> pd.DataFrame:
    """정렬 규칙 적용 (decide 지표 기준)."""
    if df.empty:
        return df
    dcols = [f"{l}_{decide}" for l in labels]
    if sort == "worst":
        df = df.assign(_m=df[dcols].max(axis=1)).sort_values(
            "_m", ascending=False).drop(columns="_m")
    elif sort == "delta" and two and delta_col:
        df = df.assign(_d=df[delta_col].abs()).sort_values(
            "_d", ascending=False, na_position="last").drop(columns="_d")
    else:                                                     # key
        df = df.sort_values("key")
    return df.reset_index(drop=True)


# ───────────────────────────────────────────────────────────────────────
# 5. 단어단위 오류분석 (치환/삽입/삭제)
# ───────────────────────────────────────────────────────────────────────
def build_error_analysis(
    data: dict[str, dict],
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """한 모델의 단어단위 오류를 정답↔예측 정렬로 집계.

    Returns:
        (치환 df, 삽입 df, 삭제 df) — 각각 횟수 내림차순.
            치환: 정답단어·예측단어·횟수 / 삽입·삭제: 단어·횟수
    """
    pairs = [(s[REF_FIELD], s[HYP_FIELD]) for s in data.values()
             if s[REF_FIELD].strip()]
    if not pairs:
        return (pd.DataFrame(), pd.DataFrame(), pd.DataFrame())
    refs, hyps = zip(*pairs)

    out = jiwer.process_words(list(refs), list(hyps))
    sub: Counter = Counter()
    ins: Counter = Counter()
    dele: Counter = Counter()
    for si, chunks in enumerate(out.alignments):
        rtok, htok = out.references[si], out.hypotheses[si]
        for a in chunks:
            if a.type == "substitute":
                for r, h in zip(rtok[a.ref_start_idx:a.ref_end_idx],
                                htok[a.hyp_start_idx:a.hyp_end_idx]):
                    sub[(r, h)] += 1
            elif a.type == "insert":
                for h in htok[a.hyp_start_idx:a.hyp_end_idx]:
                    ins[h] += 1
            elif a.type == "delete":
                for r in rtok[a.ref_start_idx:a.ref_end_idx]:
                    dele[r] += 1

    sub_df = pd.DataFrame(
        [{"정답단어": r, "예측단어": h, "횟수": n}
         for (r, h), n in sub.most_common()])
    ins_df = pd.DataFrame(
        [{"단어": w, "횟수": n} for w, n in ins.most_common()])
    del_df = pd.DataFrame(
        [{"단어": w, "횟수": n} for w, n in dele.most_common()])
    return sub_df, ins_df, del_df


# ───────────────────────────────────────────────────────────────────────
# 6. 엑셀 쓰기 + 서식
# ───────────────────────────────────────────────────────────────────────
def write_excel(
    out_path: Path, summary: pd.DataFrame, sentence: pd.DataFrame,
    preds: list[dict[str, dict]], labels: list[str],
) -> None:
    """모든 시트를 한 .xlsx 로 쓰고 서식 적용."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    tags = _sheet_tags(labels)                               # 오류시트용 짧은 태그
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="요약", index=False)
        _style_header(writer.sheets["요약"])
        _autofit(writer.sheets["요약"], {"지표": 20, "해설": 55}, 24)

        sentence.to_excel(writer, sheet_name="문장별비교", index=False)
        _format_sentence(writer.sheets["문장별비교"], sentence, labels)

        for data, tag in zip(preds, tags):                  # 모델별 오류분석
            sub_df, ins_df, del_df = build_error_analysis(data)
            for kind, df in (("치환", sub_df), ("삽입", ins_df), ("삭제", del_df)):
                sheet = f"{tag}_{kind}"
                (df if not df.empty else pd.DataFrame({"(없음)": []})).to_excel(
                    writer, sheet_name=sheet, index=False)
                _style_header(writer.sheets[sheet])
                _autofit(writer.sheets[sheet], {}, 14)


def _sheet_tags(labels: list[str]) -> list[str]:
    """오류분석 시트 이름용 짧은 태그(엑셀 시트명 31자 제한 → 라벨 20자 컷 + 중복회피)."""
    tags: list[str] = []
    for lab in labels:
        t = lab[:20]
        base, i = t, 1
        while t in tags:
            t = f"{base[:18]}_{i}"
            i += 1
        tags.append(t)
    return tags


def _format_sentence(ws: Worksheet, df: pd.DataFrame, labels: list[str]) -> None:
    """문장별비교: 헤더·너비·지표 2자리·승자 초록·완전오답 연노랑·틀고정.

    승자 강조는 decide 지표(sCER 있으면 sCER, 없으면 CER) 열에만 준다.
    """
    _style_header(ws)
    ws.freeze_panes = "D2"                                   # key·duration·ref 고정
    if df.empty:
        return
    decide = "sCER" if any(f"{l}_sCER" in df.columns for l in labels) else "CER"
    metric_cols = [df.columns.get_loc(c) + 1 for c in df.columns
                   if any(str(c) == f"{l}_{m}"
                          for l in labels for m in METRIC_FNS)]
    decide_cols = [df.columns.get_loc(f"{l}_{decide}") + 1 for l in labels]
    delta_col = next((i + 1 for i, c in enumerate(df.columns)
                      if str(c).startswith("Δ")), None)

    for r in range(2, ws.max_row + 1):
        for c in metric_cols:                                # 모든 지표 2자리
            ws.cell(r, c).number_format = _FMT2
        if delta_col:
            ws.cell(r, delta_col).number_format = "+0.00;-0.00"
        vals = [(c, ws.cell(r, c).value) for c in decide_cols
                if isinstance(ws.cell(r, c).value, (int, float))]
        if len(labels) > 1 and vals:                         # 승자(최저 decide) 초록
            lo = min(v for _, v in vals)
            winners = [c for c, v in vals if v == lo]
            if len(winners) == 1:
                ws.cell(r, winners[0]).fill = _WIN_FILL
                ws.cell(r, winners[0]).font = _WIN_FONT
        for c, v in vals:                                    # 완전오답 표시
            if v >= 100:
                ws.cell(r, c).fill = _ERR_FILL

    widths = {"key": 30, "duration": 9, "ref": 34, "winner": 12}
    for lab in labels:
        widths[f"{lab}_pred"] = 34
        widths[f"{lab}_raw"] = 34
        for m in METRIC_FNS:
            widths[f"{lab}_{m}"] = 9
    _autofit(ws, widths, 16)


def _style_header(ws: Worksheet) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def _autofit(ws: Worksheet, fixed: dict[str, int], default: int) -> None:
    for idx, cell in enumerate(ws[1], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = \
            fixed.get(str(cell.value), default)


# ───────────────────────────────────────────────────────────────────────
# 7. CLI
# ───────────────────────────────────────────────────────────────────────
def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="발화 단위 결과 비교 → 엑셀 (모델 1개 이상, 같은 벤치)",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    p.add_argument("models", nargs="+", help="결과 폴더 / 모델 이름 (여러 개)")
    p.add_argument("-b", "--benchmark", required=True,
                   help="비교할 벤치마크 id (예: kiosk_cafe_order)")
    p.add_argument("-o", "--out", default=None,
                   help="출력 .xlsx (기본: BENCHMARK/results/_compare/samples_<bench>.xlsx)")
    p.add_argument("--labels", default=None,
                   help="열 이름 직접 지정 (쉼표 구분, 모델 개수와 일치)")
    p.add_argument("--metric", choices=["scer", "cer", "both"], default="both",
                   help="문장별 지표. scer=구두점+띄어쓰기 제외, cer=공백포함, "
                        "both=둘다(승자는 sCER). 기본 both")
    p.add_argument("--only-diff", action="store_true",
                   help="전원 완전정답인 발화는 제외 (판정=sCER 있으면 sCER)")
    p.add_argument("--sort", choices=["key", "worst", "delta"], default="key",
                   help="문장별비교 정렬 (key/worst/delta)")
    return p.parse_args()


# --metric 선택 → 낼 지표 목록 (표시 순서 유지).
_METRIC_CHOICE: dict[str, list[str]] = {
    "cer": ["CER"], "scer": ["sCER"], "both": ["CER", "sCER"],
}


def main() -> None:
    a = parse_args()
    model_dirs = [resolve_model_dir(m) for m in a.models]
    labels = make_labels(model_dirs, a.labels.split(",") if a.labels else None)
    preds = [load_predictions(d, a.benchmark) for d in model_dirs]
    for lab, data in zip(labels, preds):
        logger.info("Loaded predictions", model=lab, benchmark=a.benchmark,
                    utterances=len(data))

    summary = build_summary(preds, labels)
    sentence = build_sentence(preds, labels, metrics=_METRIC_CHOICE[a.metric],
                              only_diff=a.only_diff, sort=a.sort)

    out = (Path(a.out) if a.out
           else RESULTS_DIR / "_compare" / f"samples_{a.benchmark}.xlsx")
    write_excel(out, summary, sentence, preds, labels)

    logger.info("Wrote sample comparison xlsx", path=str(out),
                benchmark=a.benchmark, models=len(labels),
                rows=len(sentence))


if __name__ == "__main__":
    try:
        main()
    except (FileNotFoundError, ValueError) as e:
        logger.error("compare_samples failed", error=str(e))
        sys.exit(1)
