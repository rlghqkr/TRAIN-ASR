"""벤치마크 결과 비교 → 엑셀 (CLI).

여러 모델의 `evaluation_report.json` 을 읽어, 모델을 열(column)로 나란히 놓은
비교용 .xlsx 를 만든다. "손으로 하던 모델 비교"(모델×벤치 CER 매트릭스 +
슬라이스 + 공통셋 평균 + baseline 대비 개선폭)를 그대로 자동화한 것.

무엇을 만드나 (엑셀 시트, 앞에서부터):
    - Overview        : 모델 1행 요약 — 벤치 수, 가중/매크로 CER, baseline 대비 Δ.
                        "그래서 뭐가 제일 좋냐"를 첫 화면에서 본다.
    - CER / sCER / WER : 행=벤치마크, 열=모델. 행별 최저(=best) 셀 초록 강조.
                        맨 아래 가중평균·매크로평균 줄. baseline 주면 Δ 열 추가.
    - Slices          : 벤치마크·메타필드·값 별 CER 을 모델끼리 나란히.

비교 설계 메모(왜 이렇게):
    - 모델마다 벤치 구성이 다르면 '전체' 가중평균은 왜곡된다. 그래서 **공통셋**
      (모든 모델이 가진 벤치)만의 평균을 공정 비교 지표로 함께 낸다.
    - 큰 벤치(수천 샘플)가 가중평균을 지배하므로 **매크로(단순)평균**도 같이 본다.
    - 도메인 적응 실험(baseline vs +도메인)은 절대 CER 보다 **Δ(개선폭)** 가 핵심.

입력 1개당(=모델 1개) 다음 중 아무거나:
    - 결과 폴더            (BENCHMARK/results/<name>/  안의 evaluation_report.json 을 읽음)
    - evaluation_report.json 파일 경로
    - 모델 이름만          (BENCHMARK/results/<name>/evaluation_report.json 으로 해석)

사용법:
    # 모델 1개 (그 모델 결과만 표로)
    python scripts/compare_results.py whisper_medium_baseline_0701

    # 여러 개 나란히 비교 (폴더/이름/json 경로 섞어도 됨)
    python scripts/compare_results.py \
        whisper_medium_baseline_0701 \
        whisper_large_v3_baseline \
        -o BENCHMARK/results/_compare/medium_vs_large.xlsx

    # baseline 대비 개선폭(Δ)까지 — 도메인 적응 비교의 표준 형태
    python scripts/compare_results.py baseline_run adapt_cafe_run \
        --baseline baseline_run

    # 열 이름 직접 지정 (기본은 폴더명에서 타임스탬프만 떼서 사용)
    python scripts/compare_results.py <A> <B> --labels baseline,adapt
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from statistics import mean

import pandas as pd
import structlog
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


logger = structlog.get_logger()

# 프로젝트 루트 기준 결과 폴더 (모델 이름만 줬을 때 여기서 찾는다).
ROOT = Path(__file__).resolve().parent.parent
RESULTS_DIR = ROOT / "BENCHMARK" / "results"
REPORT_NAME = "evaluation_report.json"

# 지표 3종. (json 키, 시트 이름) — 시트 순서도 이 순서.
METRICS: list[tuple[str, str]] = [("cer", "CER"), ("scer", "sCER"), ("wer", "WER")]

# 폴더명 끝의 실행 타임스탬프( __YYMMDD_HHMMSS ) — 짧은 열이름 만들 때 떼어낸다.
_TIMESTAMP_RE = re.compile(r"__\d{6}_\d{6}$")

# Δ(개선폭) 열 이름 접두사. 이걸로 지표열/Δ열을 구분해 서식을 다르게 준다.
_DELTA_PREFIX = "Δ "

# 엑셀 스타일 상수.
_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")     # 연한 파랑 (헤더)
_BEST_FILL = PatternFill("solid", fgColor="C6EFCE")       # 연한 초록 (행 최저 CER)
_BEST_FONT = Font(color="006100")                          # 초록 글씨
_GOOD_FILL = PatternFill("solid", fgColor="C6EFCE")       # 초록 (Δ<0, 개선)
_GOOD_FONT = Font(color="006100")
_BAD_FILL = PatternFill("solid", fgColor="FFC7CE")        # 빨강 (Δ>0, 악화)
_BAD_FONT = Font(color="9C0006")
_FOOTER_FILL = PatternFill("solid", fgColor="F2F2F2")     # 회색 (평균 줄)
_MISSING = None  # 결측(그 모델엔 없는 벤치) = 빈 칸

# 엑셀 표시 포맷.
_FMT_METRIC = "0.00"        # 지표(2자리)
_FMT_DELTA = "+0.00;-0.00"  # Δ (부호 항상 표시)
_FMT_INT = "#,##0"          # 샘플 수


# ───────────────────────────────────────────────────────────────────────
# 1. 입력 해석 & 로딩
# ───────────────────────────────────────────────────────────────────────
def resolve_report(arg: str) -> Path:
    """CLI 인자 하나 → evaluation_report.json 경로.

    폴더 / json 파일 / 모델 이름 순으로 시도하고, 다 실패하면 명시적 에러.

    Args:
        arg: 결과 폴더, json 파일 경로, 또는 모델 이름.

    Returns:
        존재가 확인된 evaluation_report.json 경로.

    Raises:
        FileNotFoundError: 어떤 방식으로도 리포트를 못 찾은 경우 (시도한 경로 나열).
    """
    p = Path(arg)
    tried: list[Path] = []

    # 1) json 파일 직접 지정
    if p.suffix == ".json":
        tried.append(p)
        if p.is_file():
            return p

    # 2) 결과 폴더 (폴더/evaluation_report.json)
    cand = p / REPORT_NAME
    tried.append(cand)
    if cand.is_file():
        return cand

    # 3) 모델 이름 (RESULTS_DIR/<name>/evaluation_report.json)
    cand = RESULTS_DIR / arg / REPORT_NAME
    tried.append(cand)
    if cand.is_file():
        return cand

    raise FileNotFoundError(
        f"'{arg}' 에서 {REPORT_NAME} 를 못 찾음. 시도한 경로:\n  "
        + "\n  ".join(str(t) for t in tried)
    )


def load_reports(args: list[str]) -> list[dict]:
    """인자 목록 → 리포트 dict 목록 (원본 JSON 그대로 + '_path' 부가)."""
    reports: list[dict] = []
    for a in args:
        path = resolve_report(a)
        data = json.loads(path.read_text(encoding="utf-8"))
        data["_path"] = path
        reports.append(data)
        logger.info("Loaded report", model=data.get("model"), path=str(path))
    return reports


def make_labels(reports: list[dict], override: list[str] | None) -> list[str]:
    """열(모델) 이름 만들기.

    기본: 폴더명에서 타임스탬프 접미사만 떼어 짧게. 단 그렇게 하면 이름이
    겹치는 경우(같은 모델 재실행 등)엔 겹치는 것들만 원래 전체 이름을 쓴다.
    --labels 로 직접 주면 그대로 사용.
    """
    if override is not None:
        if len(override) != len(reports):
            raise ValueError(
                f"--labels 개수({len(override)})가 모델 개수({len(reports)})와 다름"
            )
        return override

    full = [r["_path"].parent.name for r in reports]      # 폴더명 = 전체 이름
    short = [_TIMESTAMP_RE.sub("", name) for name in full]  # 타임스탬프 뗀 짧은 이름

    # 짧게 줄였더니 겹치는 이름은 → 그 그룹만 전체 이름으로 되돌린다(구분 위해).
    dupes = {s for s in short if short.count(s) > 1}
    return [full[i] if short[i] in dupes else short[i] for i in range(len(reports))]


def resolve_baseline(
    baseline: str | None, labels: list[str], reports: list[dict],
) -> str | None:
    """--baseline 인자를 실제 열(라벨)로 확정.

    라벨/폴더명/JSON model 필드 어느 것으로 줘도 찾도록 관대하게 매칭한다
    (--labels 로 이름을 바꿔도 원래 폴더명으로 baseline 지정 가능). 완전일치를
    먼저 시도하고, 없으면 부분일치. 여러 열에 걸리면 명시적 에러(모호함 금지).

    Returns:
        확정된 라벨(열 이름). baseline 미지정이면 None.
    """
    if baseline is None:
        return None

    # 열마다 매칭 후보(별칭) 모음: 라벨 + 폴더명 + JSON model 필드.
    aliases: list[set[str]] = [
        {lab, r["_path"].parent.name, str(r.get("model", ""))}
        for lab, r in zip(labels, reports)
    ]
    exact = [i for i, al in enumerate(aliases) if baseline in al]
    if len(exact) == 1:
        return labels[exact[0]]
    partial = [i for i, al in enumerate(aliases)
               if any(baseline in a for a in al)]
    if len(partial) == 1:
        return labels[partial[0]]

    raise ValueError(
        f"--baseline '{baseline}' 를 열에서 특정 못함(0개 또는 중복). 후보 열: {labels}"
    )


# ───────────────────────────────────────────────────────────────────────
# 2. 데이터 정합성 점검 (Data 팀 관점)
# ───────────────────────────────────────────────────────────────────────
def check_sample_consistency(reports: list[dict], labels: list[str]) -> None:
    """같은 벤치인데 모델마다 샘플 수가 다르면 경고.

    샘플 수가 다르면 CER 비교가 사과-오렌지가 되므로 조용히 넘기지 않고
    WARNING 을 남긴다(값은 그대로 표에 실어 사용자가 판단하게 함).
    """
    for bid in ordered_benchmarks(reports):
        counts = {
            lab: r["benchmarks"][bid]["samples"]
            for lab, r in zip(labels, reports)
            if bid in r["benchmarks"]
        }
        if len(set(counts.values())) > 1:
            logger.warning(
                "Sample-count mismatch — CER 비교 주의",
                benchmark=bid, samples=counts,
            )


# ───────────────────────────────────────────────────────────────────────
# 3. 벤치마크 지표 매트릭스 (CER/sCER/WER 시트)
# ───────────────────────────────────────────────────────────────────────
def ordered_benchmarks(reports: list[dict]) -> list[str]:
    """모든 모델의 벤치마크 합집합 — 처음 등장한 순서 보존."""
    order: list[str] = []
    for r in reports:
        for bid in r["benchmarks"]:
            if bid not in order:
                order.append(bid)
    return order


def common_benchmarks(reports: list[dict]) -> list[str]:
    """모든 모델이 공통으로 가진 벤치마크 (등장 순서 보존)."""
    return [b for b in ordered_benchmarks(reports)
            if all(b in r["benchmarks"] for r in reports)]


def _weighted(reports: list[dict], lab_idx: int, metric: str,
              benchmarks: list[str]) -> float | None:
    """한 모델(lab_idx)의 샘플수 가중평균 (benchmarks 대상)."""
    r = reports[lab_idx]
    num = den = 0.0
    for bid in benchmarks:
        b = r["benchmarks"].get(bid)
        if b:
            num += b[metric] * b["samples"]
            den += b["samples"]
    return round(num / den, 2) if den else None


def _macro(reports: list[dict], lab_idx: int, metric: str,
           benchmarks: list[str]) -> float | None:
    """한 모델의 매크로(단순)평균 — 벤치마다 동일 가중."""
    r = reports[lab_idx]
    vals = [r["benchmarks"][bid][metric] for bid in benchmarks
            if bid in r["benchmarks"]]
    return round(mean(vals), 2) if vals else None


def build_metric_table(
    reports: list[dict], labels: list[str], metric: str,
    baseline: str | None,
) -> tuple[pd.DataFrame, int]:
    """한 지표(cer/scer/wer)의 비교 DataFrame + 벤치마크 행 개수.

    행=벤치마크, 열=[Benchmark, Samples, <model...>, (Best), (Δ...)].
    맨 아래에 가중평균·매크로평균 줄. 결측은 빈 칸(None).
    Best/공통평균/Δ 는 모델 2개 이상일 때만.
    """
    benchmarks = ordered_benchmarks(reports)
    common = common_benchmarks(reports)
    multi = len(reports) > 1
    others = [l for l in labels if l != baseline] if baseline else []

    def delta(v: float | None, base: float | None) -> float | None:
        return round(v - base, 2) if v is not None and base is not None else None

    # ── 벤치마크 행 ──
    rows: list[dict] = []
    for bid in benchmarks:
        row: dict = {"Benchmark": bid}
        vals: dict[str, float | None] = {}
        samples: int | None = None
        for lab, r in zip(labels, reports):
            b = r["benchmarks"].get(bid)
            v = round(b[metric], 2) if b else _MISSING
            vals[lab] = row[lab] = v
            if b and samples is None:
                samples = b["samples"]                        # 대표 샘플수(첫 등장)
        row["Samples"] = samples
        if multi:
            present = {l: v for l, v in vals.items() if v is not None}
            row["Best"] = min(present, key=present.get) if present else _MISSING
        for l in others:                                      # baseline 대비 Δ
            row[f"{_DELTA_PREFIX}{l}"] = delta(vals[l], vals[baseline])
        rows.append(row)

    # ── 평균 푸터 (가중=micro, 매크로=단순) ──
    # fair = Δ 를 매길 수 있는가. '전체'는 모델마다 벤치 구성이 달라(공정 X)
    # Δ 를 비우고, '공통'셋 평균에만 Δ 를 매긴다.
    footer_specs: list[tuple[str, list[str], bool]] = [
        ("가중평균 (전체)", benchmarks, False),
    ]
    if multi:
        footer_specs.append((f"가중평균 (공통 {len(common)}종)", common, True))
        footer_specs.append((f"매크로평균 (공통 {len(common)}종)", common, True))
    else:
        footer_specs.append(("매크로평균 (전체)", benchmarks, False))

    n_bench = len(rows)
    for name, benches, fair in footer_specs:
        is_macro = name.startswith("매크로")
        agg = {lab: (_macro(reports, i, metric, benches) if is_macro
                     else _weighted(reports, i, metric, benches))
               for i, lab in enumerate(labels)}
        frow: dict = {"Benchmark": name, "Samples": _MISSING, **agg}
        if multi:
            frow["Best"] = _MISSING
        for l in others:
            frow[f"{_DELTA_PREFIX}{l}"] = delta(agg[l], agg[baseline]) if fair else _MISSING
        rows.append(frow)

    cols = (["Benchmark", "Samples", *labels]
            + (["Best"] if multi else [])
            + [f"{_DELTA_PREFIX}{l}" for l in others])
    return pd.DataFrame(rows)[cols], n_bench


# ───────────────────────────────────────────────────────────────────────
# 4. Overview 요약 시트 (User 팀 관점 — 첫 화면 한눈에)
# ───────────────────────────────────────────────────────────────────────
def build_overview(
    reports: list[dict], labels: list[str], baseline: str | None,
) -> pd.DataFrame:
    """모델 1행 요약 — 벤치 수 + 가중/매크로 CER + baseline 대비 ΔCER.

    공정 비교는 '공통셋' 기준이므로 CER 가중/매크로는 공통셋으로 계산.
    참고로 '전체' 가중 CER 도 함께 싣되, 벤치 구성이 다르면 비교 불가임을 명심.
    """
    common = common_benchmarks(reports)
    all_b = ordered_benchmarks(reports)
    base_common_cer = (
        _weighted(reports, labels.index(baseline), "cer", common)
        if baseline else None
    )

    rows: list[dict] = []
    for i, (lab, r) in enumerate(zip(labels, reports)):
        cer_common = _weighted(reports, i, "cer", common)
        row = {
            "Model": lab,
            "벤치 수": len(r["benchmarks"]),
            "CER 가중(공통)": cer_common,
            "CER 매크로(공통)": _macro(reports, i, "cer", common),
            "sCER 가중(공통)": _weighted(reports, i, "scer", common),
            "WER 가중(공통)": _weighted(reports, i, "wer", common),
            "CER 가중(전체)": _weighted(reports, i, "cer", all_b),
        }
        if baseline:
            row[f"ΔCER vs {baseline}"] = (
                round(cer_common - base_common_cer, 2)
                if cer_common is not None and base_common_cer is not None
                else _MISSING
            )
        rows.append(row)
    return pd.DataFrame(rows)


# ───────────────────────────────────────────────────────────────────────
# 5. 슬라이스 매트릭스 (Slices 시트)
# ───────────────────────────────────────────────────────────────────────
def build_slice_table(reports: list[dict], labels: list[str]) -> pd.DataFrame:
    """슬라이스(벤치×필드×값) CER 을 모델끼리 나란히.

    행 정렬: 벤치마크 등장순 → 필드 등장순 → 값은 (모델 평균 CER) 높은 순
    (=취약 슬라이스가 위로). 값이 30샘플 미만이어도 전부 포함.
    """
    bench_order = ordered_benchmarks(reports)

    # (bench, field, value) 키 수집 + 대표 샘플수/모델별 CER.
    keys: list[tuple[str, str, str]] = []
    cer: dict[tuple, dict[str, float | None]] = {}
    n_of: dict[tuple, int] = {}
    for lab, r in zip(labels, reports):
        for bid, fields in r.get("slices", {}).items():
            for fname, values in fields.items():
                for val, res in values.items():
                    k = (bid, fname, val)
                    if k not in cer:
                        keys.append(k)
                        cer[k] = {}
                        n_of[k] = res["samples"]
                    cer[k][lab] = round(res["cer"], 2)

    field_order = _first_seen_order(keys, idx=1)

    def sort_key(k: tuple[str, str, str]) -> tuple:
        vals = [v for v in cer[k].values() if v is not None]
        avg = sum(vals) / len(vals) if vals else 0.0
        b_rank = bench_order.index(k[0]) if k[0] in bench_order else 1e9
        f_rank = field_order.index(k[1]) if k[1] in field_order else 1e9
        return (b_rank, f_rank, -avg)

    rows: list[dict] = []
    for k in sorted(keys, key=sort_key):
        bid, fname, val = k
        row: dict = {"Benchmark": bid, "Field": fname, "Value": val,
                     "Samples": n_of[k]}
        for lab in labels:
            row[lab] = cer[k].get(lab, _MISSING)
        rows.append(row)

    cols = ["Benchmark", "Field", "Value", "Samples", *labels]
    return pd.DataFrame(rows, columns=cols)


def _first_seen_order(keys: list[tuple], idx: int) -> list[str]:
    """키 튜플들에서 idx 번째 원소를 처음 등장 순서대로."""
    order: list[str] = []
    for k in keys:
        if k[idx] not in order:
            order.append(k[idx])
    return order


# ───────────────────────────────────────────────────────────────────────
# 6. 엑셀 쓰기 + 서식
# ───────────────────────────────────────────────────────────────────────
def write_excel(
    out_path: Path,
    overview: pd.DataFrame,
    metric_tables: dict[str, tuple[pd.DataFrame, int]],
    slice_df: pd.DataFrame,
    labels: list[str],
) -> None:
    """모든 시트를 한 .xlsx 로 쓰고 서식(강조/너비/고정)을 입힌다."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        overview.to_excel(writer, sheet_name="Overview", index=False)
        _format_overview_sheet(writer.sheets["Overview"], overview, labels)

        for sheet, (df, n_bench) in metric_tables.items():
            df.to_excel(writer, sheet_name=sheet, index=False)
            _format_metric_sheet(writer.sheets[sheet], df, labels, n_bench)

        slice_df.to_excel(writer, sheet_name="Slices", index=False)
        _format_slice_sheet(writer.sheets["Slices"], slice_df, labels)


def _format_overview_sheet(
    ws: Worksheet, df: pd.DataFrame, labels: list[str],
) -> None:
    """Overview: 헤더·너비·소수2자리·Δ 색칠(개선 초록/악화 빨강)."""
    _style_header(ws)
    ws.freeze_panes = "B2"
    delta_cols = [i + 1 for i, c in enumerate(df.columns) if c.startswith("Δ")]
    num_cols = [i + 1 for i, c in enumerate(df.columns) if c != "Model"]

    for r in range(2, ws.max_row + 1):
        for c in num_cols:
            cell = ws.cell(r, c)
            if c in delta_cols:
                cell.number_format = _FMT_DELTA
                _paint_delta(cell)
            elif "벤치 수" in str(df.columns[c - 1]):
                cell.number_format = _FMT_INT
            else:
                cell.number_format = _FMT_METRIC

    widths = {"Model": 42, "벤치 수": 8}
    for idx, cell in enumerate(ws[1], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = \
            widths.get(str(cell.value), 16)


def _format_metric_sheet(
    ws: Worksheet, df: pd.DataFrame, labels: list[str], n_bench: int,
) -> None:
    """CER/sCER/WER 시트: 헤더·너비·소수2자리·행최저 초록·Δ 색칠·푸터 회색·틀고정."""
    _style_header(ws)
    ws.freeze_panes = "C2"                                   # 헤더행 + 왼쪽 2열 고정

    label_cols = [df.columns.get_loc(l) + 1 for l in labels]
    delta_cols = [i + 1 for i, c in enumerate(df.columns)
                  if str(c).startswith(_DELTA_PREFIX)]
    samples_col = df.columns.get_loc("Samples") + 1

    for r in range(2, ws.max_row + 1):
        is_footer = r > 1 + n_bench                          # 평균 줄인가
        for c in label_cols:
            ws.cell(r, c).number_format = _FMT_METRIC
        for c in delta_cols:
            ws.cell(r, c).number_format = _FMT_DELTA
            _paint_delta(ws.cell(r, c))
        ws.cell(r, samples_col).number_format = _FMT_INT

        if is_footer:
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).fill = _FOOTER_FILL
                ws.cell(r, c).font = Font(bold=True)
            for c in delta_cols:                             # Δ 색은 유지
                _paint_delta(ws.cell(r, c))
            continue

        if len(labels) > 1:                                  # 행 최저 CER 강조
            _highlight_row_min(ws, r, label_cols)

    _set_widths(ws, {"Benchmark": 46, "Samples": 10, "Best": 22}, labels, 17)


def _format_slice_sheet(ws: Worksheet, df: pd.DataFrame, labels: list[str]) -> None:
    """Slices 시트: 헤더·너비·소수2자리·행최저 강조·틀고정."""
    _style_header(ws)
    ws.freeze_panes = "E2"                                   # 키 4열 고정

    label_cols = [df.columns.get_loc(l) + 1 for l in labels]
    samples_col = df.columns.get_loc("Samples") + 1

    for r in range(2, ws.max_row + 1):
        for c in label_cols:
            ws.cell(r, c).number_format = _FMT_METRIC
        ws.cell(r, samples_col).number_format = _FMT_INT
        if len(labels) > 1:
            _highlight_row_min(ws, r, label_cols)

    _set_widths(ws, {"Benchmark": 30, "Field": 10, "Value": 14, "Samples": 10},
                labels, 17)


def _highlight_row_min(ws: Worksheet, r: int, label_cols: list[int]) -> None:
    """그 행에서 CER 최저(=best) 셀을 초록으로."""
    vals = [(c, ws.cell(r, c).value) for c in label_cols
            if isinstance(ws.cell(r, c).value, (int, float))]
    if vals:
        best_c = min(vals, key=lambda cv: cv[1])[0]
        ws.cell(r, best_c).fill = _BEST_FILL
        ws.cell(r, best_c).font = _BEST_FONT


def _paint_delta(cell) -> None:
    """Δ 셀: 음수(개선) 초록 / 양수(악화) 빨강 / 0·빈칸 그대로."""
    v = cell.value
    if not isinstance(v, (int, float)):
        return
    if v < 0:
        cell.fill, cell.font = _GOOD_FILL, _GOOD_FONT
    elif v > 0:
        cell.fill, cell.font = _BAD_FILL, _BAD_FONT


def _style_header(ws: Worksheet) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def _set_widths(ws: Worksheet, fixed: dict[str, int],
                labels: list[str], label_w: int) -> None:
    """헤더 이름 기준으로 열 너비 지정 (fixed 에 없으면 모델/Δ열=label_w)."""
    for idx, cell in enumerate(ws[1], start=1):
        name = str(cell.value)
        is_model = name in labels or name.startswith(_DELTA_PREFIX)
        width = fixed.get(name, label_w if is_model else 14)
        ws.column_dimensions[get_column_letter(idx)].width = width


# ───────────────────────────────────────────────────────────────────────
# 7. CLI
# ───────────────────────────────────────────────────────────────────────
def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="벤치마크 결과 비교 → 엑셀 (모델 1개 이상)",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    p.add_argument("models", nargs="+",
                   help="결과 폴더 / evaluation_report.json / 모델 이름 (여러 개 가능)")
    p.add_argument("-o", "--out", default=None,
                   help="출력 .xlsx 경로 (기본: BENCHMARK/results/_compare/compare_<n>models.xlsx)")
    p.add_argument("--baseline", default=None,
                   help="기준 모델(라벨/부분일치). 주면 나머지 모델에 Δ(개선폭) 열 추가")
    p.add_argument("--labels", default=None,
                   help="열 이름 직접 지정 (쉼표 구분, 모델 개수와 일치)")
    return p.parse_args()


def main() -> None:
    a = parse_args()
    reports = load_reports(a.models)
    labels = make_labels(reports, a.labels.split(",") if a.labels else None)
    baseline = resolve_baseline(a.baseline, labels, reports)

    check_sample_consistency(reports, labels)               # Data 정합성 경고

    overview = build_overview(reports, labels, baseline)
    metric_tables = {
        sheet: build_metric_table(reports, labels, key, baseline)
        for key, sheet in METRICS
    }
    slice_df = build_slice_table(reports, labels)

    out = (Path(a.out) if a.out
           else RESULTS_DIR / "_compare" / f"compare_{len(reports)}models.xlsx")
    write_excel(out, overview, metric_tables, slice_df, labels)

    # 콘솔 요약 — 공정 비교 지표(공통셋 CER 가중평균)를 한 줄로.
    summary = dict(zip(overview["Model"], overview["CER 가중(공통)"]))
    logger.info("Wrote comparison xlsx", path=str(out),
                models=len(reports), benchmarks=metric_tables["CER"][1],
                baseline=baseline, cer_weighted_common=summary)


if __name__ == "__main__":
    try:
        main()
    except (FileNotFoundError, ValueError) as e:
        logger.error("compare_results failed", error=str(e))
        sys.exit(1)
